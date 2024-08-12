import os
import docx
import zipfile

import json
import locale

import openpyxl


from datetime import datetime as dt
from docxtpl import DocxTemplate

from project.settings import MEDIA_ROOT, MEDIA_URL

locale.setlocale(locale.LC_ALL, '')

def filling_excel():
    wb = openpyxl.load_workbook(filename='t.xlsx')
    sheet = wb['q']

    doc = DocxTemplate('template2.docx')

    for num in range(2,len(list(sheet.rows))+1):
        
        name = sheet['B'+str(num)].value
        last_name = sheet['A'+str(num)].value
        company = sheet['D'+str(num)].value
        start_data = sheet['F'+str(num)].value
        #end_data = sheet['G'+str(num)].value.date()
        end_data = sheet['G'+str(num)].value
        
        context = {
        'post': company,
        'first_middle': name,
        'fio': last_name,
        'date': start_data,
        'contract_date': end_data  
        }
        
        doc.render(context)
        doc.save(last_name+' заявление на отпуск.docx')

def filling_doc():
    locale.setlocale(locale.LC_ALL, '')
    doc = DocxTemplate("template.docx")
    user = json.load(open(os.path.join(os.getcwd(), 'user_info.json'), encoding='utf-8'))
    if not os.path.isdir(os.path.join(os.getcwd(), 'personal')):
        os.mkdir(os.path.join(os.getcwd(), 'personal'))
    for usr in user:
        print(f'\r[+] Заполняю: {usr["last_name"]}', end='')
        data = {'manager': 'И. С. Иванов', 'reason': 'реструктуризацией и оптимизацией закваски',
                'date': dt.strftime(dt.now(), '%d %B %Y'), 'fio': f'{usr["last_name"]} {usr["first_name"]} '
                                                                  f'{usr["middle_name"]}', 'post': usr['post'],
                'first_middle': f'{usr["first_name"]} {usr["middle_name"]}',
                'contract_date': dt.strftime(dt.strptime(usr['contract_date'], "%d.%m.%Y"), '%d %B %Y'),
                'contract_num': usr['contract_num'], 'day_x': '01.07.2021'}
        doc.render(data)
        doc.save(os.path.join(os.getcwd(), 'personal',
                              f'{usr["last_name"]} {usr["first_name"]} {usr["middle_name"]}.docx'))


# Import images from docx
def import_images(doc):
    img_dir = os.path.join(MEDIA_ROOT, 'images')

    # Create directory to save all image files if it doesn't exist
    if not os.path.exists(img_dir):
        os.makedirs(img_dir)
    
    # Extract all images from .docx
    with zipfile.ZipFile(doc.file, 'r') as zipFile:
        filelist = zipFile.namelist()
        for filename in filelist:
            if filename.startswith('word/media/'):
                zipFile.extract(filename, path=img_dir)

    return img_dir

# Check for .emf images format and convert them and
# save all 'rId:filenames' relationships in an dictionary named rels
def relate_images(img_dir, doc_file):
    rels = {}
    for r in doc_file.part.rels.values():
        if isinstance(r._target, docx.parts.image.ImagePart):
            img = os.path.basename(r._target.partname)
            if img.split('.')[-1] == 'emf':
                dir = os.path.join(img_dir, 'word/media')
                command = 'inkscape --file {0} --export-plain-svg {1}.svg'.format(os.path.join(dir, img), os.path.join(dir, img.split('.')[0]))
                if os.system(command) != 0:
                    print('Could not import .docx images properly. Please, install inkscape. \'$ apt install inkscape\'')
                img = img.split('.')[0] + '.svg'
            rels[r.rId] = img

    return rels


def import_docx(Model, doc):
    # Declare variables for text extraction
    doc_file = docx.Document(doc.file)
    obj = Model()
    text = ''

    # Import and relate all images
    img_dir = import_images(doc)
    rels = relate_images(img_dir, doc_file)

    # Iterate over document paragraphs
    for paragraph in doc_file.paragraphs:
        # If heading paragraph then create a new chapter
        if paragraph.style.name.split(' ')[0] == 'Heading':
            # If chapter is not empty, save it
            if obj.title:
                obj.text = text
                obj.save()
            obj = Model.objects.create(title=paragraph.text.strip(), document=doc)
            text = ''
        # If paragraph has an image, insert an image tag with the image file
        elif 'Graphic' in paragraph._p.xml:
            for rId in rels:
                if rId in paragraph._p.xml:
                    text += ('\n<img style="width: 50vw;" src="' + os.path.join(MEDIA_URL, 'images/word/media', rels[rId]) + '">')
        # If paragraph has text, just insert text inside paragraph tags
        else:
            text += ('\n<p class="paragraph">' + paragraph.text + '</p>')
    # Save the remaining object
    obj.text = text
    obj.save()

# Delete .docx document and it's image folder
def delete_docx(doc):
    if doc.file:
        if os.path.exists(doc.file.path):
            os.remove(doc.file.path)
    img_dir = os.path.join(MEDIA_ROOT, 'images')
    for root, dirs, files in os.walk(img_dir, topdown=False):
        for name in files:
            os.remove(os.path.join(root, name))
        for name in dirs:
            os.rmdir(os.path.join(root, name))

